"""Offline independent checks for canonical native reconstruction acceptance."""
from __future__ import annotations

import base64
import json
import tempfile
import unittest
import urllib.parse
import zlib
from pathlib import Path

from native_qa import NativeAuditError, audit_native, compare_reference, render_pptx


def fixture():
    return {"schema_version": "1.0.0", "figure_id": "sample", "kind": "diagram",
            "nodes": [{"id": "a", "label": "Input"}, {"id": "b", "label": "Encoder"}],
            "edges": [{"id": "ab", "source": "a", "target": "b", "directed": True}],
            "groups": [], "labels": [], "constraints": {"required_labels": ["Encoder"]}}


class NativeAcceptanceTests(unittest.TestCase):
    def setUp(self):
        self.t = tempfile.TemporaryDirectory()
        self.addCleanup(self.t.cleanup)
        self.root = Path(self.t.name)

    def test_native_pptx_with_source_text_stays_visual_pending(self):
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
        from pptx.util import Inches
        pptx = Presentation()
        slide = pptx.slides.add_slide(pptx.slide_layouts[6])
        for idx, (name, label) in enumerate((("a", "Input"), ("b", "Encoder"))):
            s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1 + idx * 3), Inches(2), Inches(2), Inches(1))
            s.text = label
            s.name = "rfs-node:" + name
        c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(3), Inches(2.5), Inches(4), Inches(2.5))
        c.name = "connector-untagged"
        dest = self.root / "sample.pptx"
        pptx.save(dest)
        result = audit_native(fixture(), dest)
        self.assertEqual(result["structure_check"], "PENDING_UNTAGGED_OBJECTS")
        self.assertEqual(result["missing_native_labels"], [])
        self.assertTrue(result["native_source_editable_elements"])
        self.assertEqual(result["reference_comparison"], "PENDING")

    def test_tagged_pptx_native_connections_match_scientific_graph(self):
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
        from pptx.util import Inches
        from pptx.oxml.xmlchemy import OxmlElement
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        nodes = []
        for index, (identity, label) in enumerate((("a", "Input"), ("b", "Encoder"))):
            native = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                            Inches(1 + index * 3), Inches(2), Inches(2), Inches(1))
            native.name = "rfs-node:" + identity
            native.text = label
            nodes.append(native)
        edge = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(3), Inches(2.5),
                                          Inches(4), Inches(2.5))
        edge.name = "rfs-edge:ab"
        edge.begin_connect(nodes[0], 3)
        edge.end_connect(nodes[1], 1)
        element = OxmlElement("a:tailEnd")
        element.set("type", "triangle")
        edge._element.spPr.get_or_add_ln().append(element)
        path = self.root / "tagged.pptx"
        presentation.save(path)
        result = audit_native(fixture(), path)
        self.assertEqual(result["structure_check"], "PASS")
        self.assertTrue(result["object_tags_complete"])
        self.assertEqual(result["target_render"], "NOT_RUN")

    def test_tagged_pptx_reverse_arrow_is_rejected(self):
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
        from pptx.util import Inches
        from pptx.oxml.xmlchemy import OxmlElement
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        objects = []
        for index, (identity, label) in enumerate((("a", "Input"), ("b", "Encoder"))):
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1 + index * 3),
                                            Inches(2), Inches(2), Inches(1))
            shape.name, shape.text = "rfs-node:" + identity, label
            objects.append(shape)
        edge = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(3), Inches(2.5),
                                          Inches(4), Inches(2.5))
        edge.name = "rfs-edge:ab"
        edge.begin_connect(objects[0], 3)
        edge.end_connect(objects[1], 1)
        # a:headEnd means BEGINNING, so the diagram incorrectly points to Input.
        reverse = OxmlElement("a:headEnd")
        reverse.set("type", "triangle")
        edge._element.spPr.get_or_add_ln().append(reverse)
        path = self.root / "backwards.pptx"
        presentation.save(path)
        outcome = audit_native(fixture(), path)
        self.assertEqual(outcome["status"], "FAIL")
        self.assertIn("Tagged connector arrow direction is wrong or ambiguous: ab", outcome["failures"])

    def test_flattened_pptx_fails(self):
        from pptx import Presentation
        from PIL import Image
        from io import BytesIO
        buf = BytesIO()
        Image.new("RGB", (240, 140), "white").save(buf, "PNG")
        pptx = Presentation()
        slide = pptx.slides.add_slide(pptx.slide_layouts[6])
        slide.shapes.add_picture(BytesIO(buf.getvalue()), 0, 0)
        path = self.root / "flattened.pptx"
        pptx.save(path)
        result = audit_native(fixture(), path)
        self.assertEqual(result["status"], "FAIL")
        self.assertIn("Input", result["missing_native_labels"])

    def test_drawio_tagged_exact_topology(self):
        xml = '''<mxfile><diagram id="one"><mxGraphModel><root><mxCell id="0"/><mxCell id="1" parent="0"/>
        <mxCell id="rfs-node:a" value="Input" vertex="1" parent="1"/>
        <mxCell id="rfs-node:b" value="Encoder" vertex="1" parent="1"/>
        <mxCell id="rfs-edge:ab" edge="1" source="rfs-node:a" target="rfs-node:b" style="endArrow=classic;" parent="1"/>
        </root></mxGraphModel></diagram></mxfile>'''
        path = self.root / "sample.drawio"
        path.write_text(xml)
        report = audit_native(fixture(), path)
        self.assertEqual(report["structure_check"], "PASS")
        self.assertTrue(report["object_tags_complete"])
        self.assertEqual(report["target_render"], "NOT_RUN")

    def test_drawio_tagged_swapped_labels_are_rejected(self):
        source = '''<mxGraphModel><root><mxCell id="0"/><mxCell id="1" parent="0"/>
        <mxCell id="rfs-node:a" value="Encoder" vertex="1" parent="1"/>
        <mxCell id="rfs-node:b" value="Input" vertex="1" parent="1"/>
        <mxCell id="rfs-edge:ab" edge="1" source="rfs-node:a" target="rfs-node:b" style="endArrow=classic;" parent="1"/>
        </root></mxGraphModel>'''
        path = self.root / "swapped.drawio"
        path.write_text(source)
        result = audit_native(fixture(), path)
        self.assertEqual(result["status"], "FAIL")
        self.assertIn("Tagged node label mismatch: a", result["failures"])

    def test_drawio_reversed_edge_fails(self):
        xml = '''<mxGraphModel><root><mxCell id="0"/><mxCell id="1" parent="0"/>
          <mxCell id="rfs-node:a" value="Input" vertex="1" parent="1"/>
          <mxCell id="rfs-node:b" value="Encoder" vertex="1" parent="1"/>
          <mxCell id="rfs-edge:ab" edge="1" source="rfs-node:b" target="rfs-node:a" style="endArrow=classic;" parent="1"/>
          </root></mxGraphModel>'''
        path = self.root / "reverse.drawio"
        path.write_text(xml)
        report = audit_native(fixture(), path)
        self.assertEqual(report["status"], "FAIL")
        self.assertIn("Tagged connector endpoint mismatch: ab", report["failures"])

    def test_drawio_compressed_xml(self):
        inner = '''<mxGraphModel><root><mxCell id="0"/><mxCell id="1" parent="0"/>
          <mxCell id="rfs-node:a" value="Input" vertex="1" parent="1"/>
          <mxCell id="rfs-node:b" value="Encoder" vertex="1" parent="1"/>
          <mxCell id="rfs-edge:ab" edge="1" source="rfs-node:a" target="rfs-node:b" style="endArrow=classic;" parent="1"/>
          </root></mxGraphModel>'''
        compressor = zlib.compressobj(wbits=-zlib.MAX_WBITS)
        pack = compressor.compress(urllib.parse.quote(inner).encode()) + compressor.flush()
        blob = base64.b64encode(pack).decode()
        path = self.root / "compressed.drawio"
        path.write_text(f'<mxfile><diagram>{blob}</diagram></mxfile>')
        self.assertEqual(audit_native(fixture(), path)["structure_check"], "PASS")

    def test_oversize_and_doctype_fail(self):
        path = self.root / "bad.drawio"
        path.write_bytes(b'<!DOCTYPE foo [<!ENTITY x "hello">]><mxGraphModel/>')
        with self.assertRaises(NativeAuditError):
            audit_native(fixture(), path)

    def test_symlink_path_is_rejected(self):
        real = self.root / "plain.drawio"
        real.write_text("<mxGraphModel/>")
        link = self.root / "alias.drawio"
        link.symlink_to(real)
        with self.assertRaises(NativeAuditError):
            audit_native(fixture(), link)

    @unittest.skipUnless(__import__('shutil').which('libreoffice'), "LibreOffice unavailable")
    def test_saved_pptx_opens_in_real_renderer(self):
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
        from pptx.util import Inches
        p = Presentation(); s = p.slides.add_slide(p.slide_layouts[6]);
        for idx, label in enumerate(["Input", "Encoder"]):
            t = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1+idx*3), Inches(2), Inches(2), Inches(1))
            t.text = label
        s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(3), Inches(2.5), Inches(4), Inches(2.5))
        file = self.root / "real.pptx"; p.save(file)
        img = render_pptx(file)
        self.assertTrue(img.startswith(b'\x89PNG'))
        preview = self.root / 'actual.png'
        preview.write_bytes(img)
        report = compare_reference(preview, img)
        self.assertEqual(report['normalized_mean_absolute_pixel_difference'], 0.0)
        self.assertIn('NOT_HIGH_FIDELITY_PROOF', report['interpretation'])

if __name__ == '__main__': unittest.main()
